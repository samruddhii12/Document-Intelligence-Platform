"""Deterministic bounded dataset analysis. AI supplies plans, never executable code."""
import csv
import gzip
import io
import json
import threading
from typing import Literal, Any
import numpy as np
import pandas as pd
import duckdb
from pydantic import BaseModel, Field, ConfigDict
from backend.config import settings
from backend.services.text_extractor import check_archive

class Filter(BaseModel):
    model_config=ConfigDict(extra="forbid")
    column: str
    operator: Literal["eq","ne","gt","ge","lt","le","contains","is_null","not_null","in"]
    value: Any = None

class Measure(BaseModel):
    model_config=ConfigDict(extra="forbid")
    column: str | None = None
    function: Literal["sum","mean","min","max","count","median"]

class AnalysisPlan(BaseModel):
    model_config=ConfigDict(extra="forbid")
    sheet: str | None = None
    operation: Literal["preview","describe","missing","duplicates","aggregate","correlation","anomalies","trend"] = "preview"
    columns: list[str] = Field(default_factory=list,max_length=20)
    filters: list[Filter] = Field(default_factory=list,max_length=10)
    group_by: list[str] = Field(default_factory=list,max_length=5)
    measures: list[Measure] = Field(default_factory=list,max_length=10)
    date_column: str | None = None
    interval: Literal["day","week","month","year"] = "month"
    sort_by: str | None = None
    descending: bool = False
    limit: int = Field(default=100,ge=1,le=500)
    chart: Literal["none","bar","line","scatter","histogram"] = "none"
    chart_x: str | None = None
    chart_y: str | None = None

class DatasetQuestion(BaseModel):
    question: str = Field(min_length=1,max_length=4000)
    sheet: str | None = None


def safe_json(value):
    if isinstance(value,pd.DataFrame):
        return json.loads(value.to_json(orient="records",date_format="iso"))
    return json.loads(json.dumps(value,default=str,allow_nan=False))

def canonical_frame(raw,headers):
    raw.columns=[f"c{index+1}" for index in range(len(raw.columns))]
    mapping=[{"id":column,"name":str(headers[index]) if index<len(headers) and headers[index] is not None else f"Column {index+1}"}
             for index,column in enumerate(raw.columns)]
    return raw,mapping

def parse_dataset(content,extension):
    tables={}
    if extension=="csv":
        text=content.decode("utf-8-sig")
        if "\x00" in text: raise ValueError("CSV contains binary content")
        try: delimiter=csv.Sniffer().sniff(text[:8192],delimiters=",;\t|").delimiter
        except csv.Error: delimiter=","  # Single-column CSV is valid.
        headers=next(csv.reader(io.StringIO(text),delimiter=delimiter),[])
        if not headers or len(headers)>settings.MAX_DATASET_COLUMNS: raise ValueError("CSV has too many columns or no header")
        raw=pd.read_csv(io.StringIO(text),sep=delimiter,nrows=settings.MAX_DATASET_ROWS+1)
        if len(raw)>settings.MAX_DATASET_ROWS: raise ValueError("Dataset exceeds the row limit")
        frame,mapping=canonical_frame(raw,headers); tables["CSV"]=(frame,mapping)
    elif extension=="xlsx":
        check_archive(content)
        from openpyxl import load_workbook
        workbook=load_workbook(io.BytesIO(content),read_only=True,data_only=True,keep_links=False)
        try:
            if len(workbook.worksheets)>settings.MAX_XLSX_SHEETS: raise ValueError("Workbook has too many sheets")
            total_rows=0
            for sheet in workbook.worksheets:
                if sheet.max_column>settings.MAX_DATASET_COLUMNS or sheet.max_row>settings.MAX_DATASET_ROWS+1:
                    raise ValueError("Worksheet exceeds row/column limits")
                rows=sheet.iter_rows(values_only=True)
                headers=next(rows,())
                if not headers: continue
                values=[]
                for row in rows:
                    values.append(row); total_rows+=1
                    if total_rows>settings.MAX_DATASET_ROWS: raise ValueError("Workbook exceeds the total row limit")
                raw=pd.DataFrame(values,columns=[f"c{i}" for i in range(len(headers))])
                frame,mapping=canonical_frame(raw,headers); tables[sheet.title]=(frame,mapping)
        finally: workbook.close()
    else: raise ValueError("Only CSV/XLSX datasets are supported")
    if not tables: raise ValueError("Dataset has no tables")
    total_memory=0
    for frame,mapping in tables.values():
        if len(frame.columns)>settings.MAX_DATASET_COLUMNS: raise ValueError("Dataset exceeds column limit")
        total_memory+=int(frame.memory_usage(deep=True).sum())
    if total_memory>settings.MAX_UNPACKED_MB*1024*1024: raise ValueError("Dataset exceeds the in-memory size limit")
    return tables

def profile_dataset(tables):
    result={"sheets":{},"version":1,"formula_policy":"Formula cached values only; formulas and macros are never executed."}
    for name,(frame,mapping) in tables.items():
        numeric=frame.select_dtypes(include="number")
        result["sheets"][name]={"rows":len(frame),"columns":[{**info,"dtype":str(frame[info["id"]].dtype),
             "missing":int(frame[info["id"]].isna().sum()),"unique":int(frame[info["id"]].nunique())} for info in mapping],
             "duplicate_rows":int(frame.duplicated().sum()),"preview":safe_json(frame.head(20)),
             "statistics":safe_json(numeric.describe().reset_index()) if len(numeric.columns) else [],
             "scope":"full dataset"}
    return result

def serialize_tables(tables):
    data={name:{"table":json.loads(frame.to_json(orient="table",date_format="iso",index=False)),"columns":columns}
          for name,(frame,columns) in tables.items()}
    encoded=json.dumps(data,allow_nan=False).encode()
    if len(encoded)>settings.MAX_UNPACKED_MB*1024*1024: raise ValueError("Normalized dataset exceeds the size limit")
    return gzip.compress(encoded)

def deserialize_tables(content):
    # Content is generated by our parser, never an uploaded pickle.
    with gzip.GzipFile(fileobj=io.BytesIO(content)) as stream:
        encoded=stream.read(settings.MAX_UNPACKED_MB*1024*1024+1)
    if len(encoded)>settings.MAX_UNPACKED_MB*1024*1024: raise ValueError("Normalized dataset exceeds the size limit")
    data=json.loads(encoded)
    return {name:(pd.read_json(io.StringIO(json.dumps(info["table"])),orient="table"),info["columns"]) for name,info in data.items()}

def identifier(name,columns):
    if name not in columns: raise ValueError(f"Unknown column: {name}")
    return '"'+name.replace('"','""')+'"'

def execute_plan(tables,plan:AnalysisPlan):
    sheet=plan.sheet or next(iter(tables))
    if sheet not in tables: raise ValueError("Unknown worksheet")
    frame,mapping=tables[sheet]; columns=set(frame.columns)
    for column in plan.columns+plan.group_by: identifier(column,columns)
    clauses=[]; parameters=[]
    for item in plan.filters:
        column=identifier(item.column,columns)
        if item.operator in {"is_null","not_null"}:
            clauses.append(column+(" IS NULL" if item.operator=="is_null" else " IS NOT NULL"))
        elif item.operator=="contains":
            if not isinstance(item.value,str) or len(item.value)>1000: raise ValueError("Contains filter requires a short string")
            clauses.append(f"strpos(CAST({column} AS VARCHAR),?)>0"); parameters.append(item.value)
        elif item.operator=="in":
            if not isinstance(item.value,list) or not 1<=len(item.value)<=100 or any(isinstance(v,(list,dict)) for v in item.value):
                raise ValueError("In filter requires 1–100 scalar values")
            clauses.append(column+" IN ("+",".join("?" for _ in item.value)+")"); parameters.extend(item.value)
        else:
            if isinstance(item.value,(list,dict)): raise ValueError("Filter value must be scalar")
            if item.value is None:
                if item.operator not in {"eq","ne"}: raise ValueError("Use null filters for missing values")
                clauses.append(column+(" IS NULL" if item.operator=="eq" else " IS NOT NULL"))
            else:
                operator={"eq":"=","ne":"!=","gt":">","ge":">=","lt":"<","le":"<="}[item.operator]
                clauses.append(f"{column} {operator} ?"); parameters.append(item.value)
    where=" WHERE "+" AND ".join(clauses) if clauses else ""
    connection=duckdb.connect(":memory:",config={"threads":1,"memory_limit":"128MB","enable_external_access":False})
    timer=threading.Timer(10,connection.interrupt); timer.daemon=True; timer.start()
    try:
        connection.register("authorized_data",frame)
        filtered=connection.execute('SELECT * FROM "authorized_data"'+where,parameters).fetchdf()
        total=len(filtered)
        operation=plan.operation
        full_result_count=None
        if operation=="preview":
            result=filtered[plan.columns or list(frame.columns)]
        elif operation=="missing":
            result=pd.DataFrame([{"column":column,"missing":int(filtered[column].isna().sum()),
                                  "percentage":float(filtered[column].isna().mean()*100) if total else 0.0} for column in plan.columns or frame.columns])
        elif operation=="duplicates":
            subset=plan.columns or None
            result=filtered[filtered.duplicated(subset=subset,keep=False)]
        elif operation=="describe":
            selected=filtered[plan.columns or list(frame.columns)]
            result=selected.describe(include="all").reset_index() if total and len(selected.columns) else pd.DataFrame()
        elif operation=="correlation":
            selected=filtered[plan.columns] if plan.columns else filtered.select_dtypes(include="number")
            if len(selected.columns)<2 or any(not pd.api.types.is_numeric_dtype(selected[c]) for c in selected):
                raise ValueError("Correlation requires at least two numeric columns")
            result=selected.corr().reset_index()
        elif operation=="anomalies":
            selected=plan.columns or list(filtered.select_dtypes(include="number").columns)
            if not selected: raise ValueError("Anomaly analysis needs numeric columns")
            records=[]; full_result_count=0
            for column in selected:
                if not pd.api.types.is_numeric_dtype(filtered[column]): raise ValueError("Anomaly analysis needs numeric columns")
                series=filtered[column]; q1,q3=series.quantile([0.25,0.75]); spread=q3-q1
                low,high=q1-1.5*spread,q3+1.5*spread
                flagged=series[(series<low)|(series>high)]
                full_result_count+=len(flagged)
                for index,value in flagged.head(max(0,plan.limit-len(records))).items():
                    records.append({"row_index":int(index),"column":column,"value":float(value),"lower_bound":float(low),"upper_bound":float(high)})
            result=pd.DataFrame(records)
        elif operation in {"aggregate","trend"}:
            if not plan.measures: raise ValueError("Choose at least one aggregation")
            groups=[identifier(column,columns) for column in plan.group_by]; expressions=[]
            if operation=="trend":
                if not plan.date_column: raise ValueError("Choose a date column for trends")
                date=identifier(plan.date_column,columns)
                groups.insert(0,f"date_trunc('{plan.interval}',TRY_CAST({date} AS TIMESTAMP))")
            for index,measure in enumerate(plan.measures):
                if measure.column is None:
                    if measure.function!="count": raise ValueError("Only count can omit its column")
                    expression="count(*)"
                else:
                    column=identifier(measure.column,columns)
                    if measure.function in {"sum","mean","median"} and not pd.api.types.is_numeric_dtype(frame[measure.column]):
                        raise ValueError("This aggregation needs a numeric column")
                    function="avg" if measure.function=="mean" else measure.function
                    expression=f"{function}({column})"
                expressions.append(expression+f' AS "measure_{index+1}"')
            select=[group+(" AS period" if index==0 and operation=="trend" else "") for index,group in enumerate(groups)]+expressions
            query='SELECT '+", ".join(select)+' FROM "authorized_data"'+where
            if groups: query+=' GROUP BY '+", ".join(groups)
            result=connection.execute(query,parameters).fetchdf()
        else: raise ValueError("Unsupported operation")
        result.columns=[str(column) for column in result.columns]
        if plan.sort_by:
            identifier(plan.sort_by,set(result.columns)); result=result.sort_values(plan.sort_by,ascending=not plan.descending)
        elif operation=="trend" and "period" in result:
            result=result.sort_values("period")
        result_count=len(result) if full_result_count is None else full_result_count
        result=result.head(plan.limit)
        chart=None
        if plan.chart!="none":
            available=set(result.columns)
            x=plan.chart_x or (str(result.columns[0]) if len(result.columns) else None)
            y=plan.chart_y or (str(result.columns[1]) if len(result.columns)>1 else None)
            identifier(x,available)
            if plan.chart!="histogram": identifier(y,available)
            chart={"type":plan.chart,"x":x,"y":y}
        return {"sheet":sheet,"operation":operation,"rows_considered":total,"result_rows":result_count,
                "truncated":result_count>plan.limit,"columns":list(result.columns),"data":safe_json(result),"chart":chart,
                "source_columns":mapping,"scope":"full filtered dataset","method":"IQR (1.5 × interquartile range)" if operation=="anomalies" else operation,
                "notes":"Correlation is not evidence of causation. Trend rows with unparseable dates appear under a null period." if operation in {"correlation","trend"} else ""}
    except duckdb.Error as error:
        raise ValueError("Analysis could not run within its limits. Check column types, filter values, or reduce the request.") from error
    finally:
        timer.cancel(); timer.join(timeout=1); connection.close()
